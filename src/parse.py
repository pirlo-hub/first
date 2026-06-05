"""把简历文件转成可发给 DeepSeek V4-Pro 的 content block 列表。

支持：
  - PDF（有文字层）→ pdfplumber 抽文字（快、省 token）
  - PDF（扫描件/图片层）→ pymupdf 转图片后以 image_url 发送
  - 图片（jpg/png/webp 等）→ base64 image_url
  - Word / txt / md → 纯文字
"""

import base64
import mimetypes
from pathlib import Path

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
TEXT_EXTS = {".txt", ".md"}


def build_content_blocks(file_path: str) -> list:
    """返回 OpenAI vision 格式的 content block 列表。"""
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext == ".pdf":
        return _pdf_blocks(path)

    if ext in IMAGE_EXTS:
        return [_image_block(path)]

    if ext == ".docx":
        return [{"type": "text", "text": _docx_text(path)}]

    if ext in TEXT_EXTS:
        return [{"type": "text", "text": path.read_text(encoding="utf-8", errors="ignore")}]

    raise ValueError(f"不支持的文件类型: {ext}（支持 pdf/docx/txt/md 及常见图片）")


def extract_text(file_path: str) -> str:
    """兼容旧调用：返回纯文字（仅 PDF/docx/txt 有意义）。"""
    blocks = build_content_blocks(file_path)
    texts = [b["text"] for b in blocks if b.get("type") == "text"]
    return "\n".join(texts)


def summarize_blocks(blocks: list) -> dict:
    """概括一组 content block，供界面核对"模型到底读到了什么"。

    返回 {"text": 拼接的文字, "n_images": 图片数量, "kind": "text"/"image"/"mixed"}。
    """
    texts = [b["text"] for b in blocks if b.get("type") == "text"]
    n_images = sum(1 for b in blocks if b.get("type") == "image_url")
    text = "\n".join(texts).strip()
    if text and n_images:
        kind = "mixed"
    elif n_images:
        kind = "image"
    else:
        kind = "text"
    return {"text": text, "n_images": n_images, "kind": kind}


# 文字层太短就认定为扫描件（按字符数）；中文简历正常都远超这个量
_MIN_PDF_TEXT_CHARS = 40


# ── 内部工具 ─────────────────────────────────────────────


def _image_block(path: Path) -> dict:
    media = mimetypes.guess_type(str(path))[0] or "image/jpeg"
    data = base64.standard_b64encode(path.read_bytes()).decode()
    return {
        "type": "image_url",
        "image_url": {"url": f"data:{media};base64,{data}"},
    }


def _pdf_blocks(path: Path) -> list:
    """先尝试 pdfplumber 抽文字（自动剔除水印）；如果是扫描件，改用 pymupdf 转图片。"""
    import pdfplumber

    pages_text = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            t = _extract_text_no_watermark(page)
            if t and t.strip():
                pages_text.append(t.strip())

    joined = "\n\n".join(pages_text).strip()
    # 文字层够多才走文字；太短（残缺/只有页眉页脚）就当扫描件转图片，
    # 避免模型拿到半截内容后"脑补"剩下的。
    if len(joined) >= _MIN_PDF_TEXT_CHARS:
        return [{"type": "text", "text": joined}]

    # 扫描件 / 文字层残缺：逐页转图片
    return _pdf_as_images(path)


def _extract_text_no_watermark(page) -> str:
    """提取一页的文字，剔除常见水印字符（旋转 / 极淡灰）。

    PDF 水印一般通过两种方式实现：
      1) 把文字旋转 45°/30° 等角度 → pdfplumber 的 char['upright'] 为 False
      2) 把字体设成接近白色的浅灰 → char['non_stroking_color'] 是高灰度值

    过滤掉这两类 char 后再排版抽文字。若过滤后剩余太少（<50% 原文），
    可能是误伤，退回原文。
    """
    try:
        original = page.extract_text() or ""

        def _keep(obj):
            if obj.get("object_type") != "char":
                return True
            # 旋转字符（水印典型特征）
            if obj.get("upright") is False:
                return False
            # 极淡灰：颜色越接近 1（白）越可能是水印背景文字
            col = obj.get("non_stroking_color")
            if isinstance(col, (int, float)) and col >= 0.75:
                return False
            if isinstance(col, (list, tuple)) and len(col) >= 3:
                r, g, b = col[0], col[1], col[2]
                if all(isinstance(x, (int, float)) for x in (r, g, b)) \
                        and r >= 0.75 and g >= 0.75 and b >= 0.75:
                    return False
            return True

        filtered = page.filter(_keep).extract_text() or ""

        # 安全网：若过滤后剩余文字少于原文 50%，可能是误伤，退回原文
        if original and len(filtered.strip()) < len(original.strip()) * 0.5:
            return original
        return filtered
    except Exception:
        # 任何异常都退回原始抽取，保证不会让水印过滤把简历搞丢
        try:
            return page.extract_text() or ""
        except Exception:
            return ""


def _pdf_as_images(path: Path) -> list:
    """用 pymupdf 把 PDF 每页渲染成 PNG，打包成 image_url block 列表。"""
    try:
        import fitz  # pymupdf
    except ImportError:
        raise RuntimeError(
            "处理图片 PDF 需要 pymupdf，请运行：pip install pymupdf"
        )

    blocks = []
    doc = fitz.open(str(path))
    for page in doc:
        pix = page.get_pixmap(dpi=200)  # 中文小字识别更稳，避免看不清而脑补
        png_bytes = pix.tobytes("png")
        data = base64.standard_b64encode(png_bytes).decode()
        blocks.append({
            "type": "image_url",
            "image_url": {"url": f"data:image/png;base64,{data}"},
        })
    doc.close()
    return blocks


def _docx_text(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)
